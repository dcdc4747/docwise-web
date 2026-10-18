"""草稿导出：Markdown 与 Word（.docx）两路出口。

口径（别自己发明）：

- **Markdown 是纯逻辑**：正文**原样**返回，末尾按引文 `order` 追加「## 参考文献」。
  每条只写 `csl_json` 里**真有的** `title` / `container-title` / `page`，
  缺哪项跳哪项，**绝不用占位串编**。
- **.docx 的重活在引擎环境**：后端只有固定的几个包，`python-docx` 不在其中，
  所以借 `DOCWISE_ENGINE_PYTHON` 指的解释器跑一段内联脚本。
  引擎侧不可用（没配 / 没装 python-docx / 超时 / 退出码非 0）→ **返回 None**，
  由路由层回 503 + 人话——**绝不返回半成品字节**。
- **有 `body_html` 就带样式**：块与字号/字体/行距/缩进/对齐由
  `services/docx_styles.py` 解析（不依赖 python-docx），再交给内联脚本落成
  run 级格式。**没 html 或解析失败** → 原样回退 `markdown_blocks`，
  **导出绝不因解析失败而失败**。
- 本模块是服务层：不碰 HTTP、不抛 `HTTPException`（见 `services/__init__.py`），
  失败一律 None，措辞交给路由层。
"""

from __future__ import annotations

import json
import logging
import os
import re
import subprocess
from collections.abc import Mapping, Sequence
from typing import Any

from ..netenv import sanitize_proxy_env
from .docx_styles import parse_body_html_checked

logger = logging.getLogger(__name__)

BIBLIOGRAPHY_HEADING = "## 参考文献"

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

# 与翻译 / 参考文献格式化共用同一个变量：后端不装 python-docx，只借引擎环境。
_ENGINE_PYTHON_ENV = "DOCWISE_ENGINE_PYTHON"

# 生成一份稿子的 docx 是"一次渲染"，与翻译（1800 秒）不是一个量级；
# 超时按失败报（None），不无声挂着。
_TIMEOUT_SECONDS = 60

# Windows 文件名保留字符 + 控制字符（含换行）一律换掉。
_UNSAFE_FILENAME_RE = re.compile(r'[\\/:*?"<>|\x00-\x1f\x7f]')
_MAX_FILENAME_STEM = 80

# Markdown 标题：`#` 到 `######` 后跟空格（`#标题` 不算，符合 CommonMark）。
_HEADING_RE = re.compile(r"^(#{1,6})[ \t]+(.*)$")


# ============================================================ 取值（ORM 行或字典）

def _field(value: object) -> str:
    """只把真有的值变成字符串：None → 空，其余 strip 后原样（**不补占位串**）。"""
    if value is None:
        return ""
    return str(value).strip()


def _attr(source: object, name: str) -> Any:
    """草稿/引文既可能是 ORM 行，也可能是字典——纯函数要好测。"""
    if isinstance(source, Mapping):
        return source.get(name)
    return getattr(source, name, None)


def _csl_of(citation: object) -> Mapping[str, Any]:
    raw = _attr(citation, "csl_json")
    return raw if isinstance(raw, Mapping) else {}


def _order_of(citation: object) -> int:
    try:
        return int(_attr(citation, "order") or 0)
    except (TypeError, ValueError):
        return 0


# ============================================================ 参考文献

def reference_line(number: int, csl: Mapping[str, Any]) -> str | None:
    """一条参考文献：`[n] 标题. 来源. 页码`，缺哪项跳哪项。

    ⚠️ **只读 CSL-JSON 真有的字段**（`title` / `container-title` / `page`），
    **绝不用"未知作者"这类占位串冒充数据**；三项都没有时这一条不出现。
    """
    parts = [
        _field(csl.get("title")),
        _field(csl.get("container-title")),
        _field(csl.get("page")),
    ]
    body = ". ".join(part for part in parts if part)
    return f"[{number}] {body}" if body else None


def reference_lines(citations: Sequence[Any]) -> list[str]:
    """按引文 `order`（文内编号）排序后逐条格式化。

    编号用 `order` 本身、不用枚举下标：正文里的 `[n]` 指的是它，
    中间缺一条也不能把后面的编号往前挤。
    """
    lines: list[str] = []
    for index, citation in enumerate(sorted(citations, key=_order_of), start=1):
        number = _order_of(citation) or index
        line = reference_line(number, _csl_of(citation))
        if line:
            lines.append(line)
    return lines


# ============================================================ Markdown

def to_markdown(draft: object, citations: Sequence[Any]) -> str:
    """正文**原样** + 末尾「## 参考文献」；没有引文时不加空标题。"""
    raw_body = _attr(draft, "body_md")
    body = "" if raw_body is None else str(raw_body)

    lines = reference_lines(citations)
    if not lines:
        return body

    section = BIBLIOGRAPHY_HEADING + "\n\n" + "\n".join(lines)
    stripped = body.rstrip("\n")
    if not stripped.strip():
        return section
    return f"{stripped}\n\n{section}"


# ============================================================ Word（.docx）

def markdown_blocks(body_md: str) -> list[dict[str, Any]]:
    """把正文切成 docx 能直接落的块（标题带级别 / 段落）。

    - `#` 到 `######` 后跟空格 → 标题（级别 1..6）；
    - 空行分段，段内单换行保留（引擎脚本落成软换行），**不丢字**；
    - 加粗 / 列表等富文本**不解释**（本期只要求标题映射），原文照落。
    """
    blocks: list[dict[str, Any]] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append({"kind": "paragraph", "text": "\n".join(paragraph)})
            paragraph.clear()

    text = (body_md or "").replace("\r\n", "\n").replace("\r", "\n")
    for raw in text.split("\n"):
        match = _HEADING_RE.match(raw)
        heading = match.group(2).strip() if match else ""
        if match and heading:
            flush()
            blocks.append(
                {"kind": "heading", "level": len(match.group(1)), "text": heading}
            )
        elif raw.strip():
            paragraph.append(raw.rstrip())
        else:
            flush()
    flush()
    return blocks


def _docx_blocks_from_html(
    parsed: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    """把 `docx_styles` 的解析结果搬成内联脚本的块（**只搬运，不改写**）。"""
    blocks: list[dict[str, Any]] = []
    for item in parsed:
        level = int(item.get("level") or 0)
        blocks.append(
            {
                "kind": "heading" if level else "paragraph",
                "level": level,
                "text": str(item.get("text") or ""),
                "bold": bool(item.get("bold")),
                "font_family": item.get("font_family"),
                "font_size_pt": item.get("font_size_pt"),
                "line_height": item.get("line_height"),
                "first_line_indent_pt": item.get("first_line_indent_pt"),
                "align": item.get("align"),
            }
        )
    return blocks


def safe_filename(title: object, suffix: str) -> str:
    """草稿标题 → 附件文件名：替换不安全字符、压掉多余空白、限长。

    Windows 保留字符与控制字符都换成 `_`；标题为空时用 `draft` 兜底
    （**不编标题**，只是个文件名）。
    """
    stem = _UNSAFE_FILENAME_RE.sub("_", _field(title))
    stem = re.sub(r"\s+", " ", stem).strip().strip(".")
    stem = stem[:_MAX_FILENAME_STEM].strip().strip(".")
    return f"{stem or 'draft'}{suffix}"


def docx_unavailable_detail() -> str:
    """路由层 503 的人话：先分清"没配解释器"，其余归为引擎侧不可用。"""
    if not _field(os.environ.get(_ENGINE_PYTHON_ENV)):
        return f"Word 导出需要引擎环境（未配置 {_ENGINE_PYTHON_ENV}）"
    return (
        "Word 导出需要引擎环境（引擎侧没有可用的 python-docx，"
        "或导出超时/进程异常）"
    )


# 内联脚本：只负责"把结构化块写成 docx"，Markdown / HTML 解析已经在上面做完。
# 用 `-c` 而不是落一个 .py 文件：不往仓库/磁盘留脚本，也避免路径推断。
# 块可带样式字段（font_family / font_size_pt / line_height /
# first_line_indent_pt / align / bold）；字段缺失就按默认样式落。
_DOCX_SCRIPT = r'''
import io
import json
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _float(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _east_asian(run, family):
    # 中文字体光设 font.name 不生效，必须同时写 eastAsia。
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), family)


def _apply_style(paragraph, block):
    family = str(block.get("font_family") or "").strip()
    size = _float(block.get("font_size_pt"))
    line = _float(block.get("line_height"))
    indent = _float(block.get("first_line_indent_pt"))
    align = _ALIGN.get(str(block.get("align") or "").strip().lower())
    fmt = paragraph.paragraph_format
    if align is not None:
        fmt.alignment = align
    if line is not None and line > 0:
        fmt.line_spacing = line
    if indent:
        fmt.first_line_indent = Pt(indent)
    bold = bool(block.get("bold"))
    for run in paragraph.runs:
        if bold:
            run.bold = True
        if family:
            run.font.name = family
            _east_asian(run, family)
        if size is not None and size > 0:
            run.font.size = Pt(size)


def _add_paragraph(doc, text):
    para = doc.add_paragraph()
    lines = text.split("\n")
    for index, line in enumerate(lines):
        run = para.add_run(line)
        if index < len(lines) - 1:
            run.add_break()
    return para


def main() -> int:
    payload = json.loads(sys.stdin.buffer.read().decode("utf-8"))

    doc = Document()
    title = str(payload.get("title") or "").strip()
    if title:
        doc.core_properties.title = title
        doc.add_heading(title, level=0)

    for block in payload.get("blocks") or []:
        text = str(block.get("text") or "")
        if not text:
            continue
        if block.get("kind") == "heading":
            level = int(block.get("level") or 1)
            para = doc.add_heading(text.replace("\n", " "), level=level)
        else:
            para = _add_paragraph(doc, text)
        _apply_style(para, block)

    references = payload.get("references") or []
    if references:
        doc.add_heading("参考文献", level=2)
        for entry in references:
            doc.add_paragraph(str(entry))

    buffer = io.BytesIO()
    doc.save(buffer)
    sys.stdout.buffer.write(buffer.getvalue())
    return 0


if __name__ == "__main__":
    sys.exit(main())
'''


def to_docx_bytes(draft: object, citations: Sequence[Any]) -> bytes | None:
    """在引擎环境里生成 .docx；**任何一步失败都返回 None**（不返回半成品）。

    失败面：没配 `DOCWISE_ENGINE_PYTHON`、解释器起不来（OSError）、
    超时、退出码非 0（最常见是没装 python-docx）、输出不是 docx 字节。
    """
    python = _field(os.environ.get(_ENGINE_PYTHON_ENV))
    if not python:
        logger.info("Word 导出跳过：未配置 %s", _ENGINE_PYTHON_ENV)
        return None

    html = _field(_attr(draft, "body_html"))
    blocks: list[dict[str, Any]] = []
    fallback = "未提供 body_html"
    if html:
        try:
            parsed, fallback = parse_body_html_checked(html)
        except Exception:
            # 双保险：解析器哪怕有 bug，也不许把导出搞挂。
            parsed, fallback = [], "HTML 解析异常"
        if parsed:
            blocks = _docx_blocks_from_html(parsed)
            fallback = ""
    if not blocks:
        if html:
            logger.info(
                "Word 导出回退 Markdown：%s", fallback or "HTML 没有可用块"
            )
        blocks = markdown_blocks(str(_attr(draft, "body_md") or ""))

    payload = {
        "title": _field(_attr(draft, "title")),
        "blocks": blocks,
        "references": reference_lines(citations),
    }

    child_env = os.environ.copy()
    # Windows 下子进程默认按 GBK 读写，中文标题/正文会被搞坏 →
    # 三个流都钉 UTF-8（与引擎适配器同一条防线）。
    child_env["PYTHONIOENCODING"] = "utf-8"
    child_env["PYTHONUTF8"] = "1"
    # 代理环境变量先洗一遍：`NO_PROXY` 里带方括号的 IPv6 会炸掉 httpx 的
    # URLPattern（引擎侧踩过，见 app/netenv.py）。
    sanitize_proxy_env(child_env)

    try:
        proc = subprocess.run(
            [python, "-c", _DOCX_SCRIPT],
            input=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            capture_output=True,
            env=child_env,
            timeout=_TIMEOUT_SECONDS,
            check=False,
        )
    except subprocess.TimeoutExpired:
        logger.warning("Word 导出超时（超过 %s 秒）", _TIMEOUT_SECONDS)
        return None
    except OSError as exc:
        logger.warning("Word 导出起不来引擎进程：%s", exc)
        return None

    if proc.returncode != 0:
        reason = (proc.stderr or b"").decode("utf-8", "replace").strip()
        logger.warning(
            "Word 导出失败（引擎退出码 %s）：%s",
            proc.returncode,
            reason[-500:] or "引擎没有给出原因",
        )
        return None

    data = proc.stdout or b""
    # docx 是 ZIP（本地文件头 `PK\x03\x04`）：没有这个前缀就不是成品。
    if not data.startswith(b"PK"):
        logger.warning("Word 导出失败：引擎没有输出 docx 字节")
        return None
    return data
